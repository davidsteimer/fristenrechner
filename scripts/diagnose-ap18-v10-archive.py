"""Read-only V0.10 archival diagnostic against a reconstructed comparator.

Replays retained V0.9/V0.10 author exports and declared patches in memory.
No XLSX, historical receipt, or trusted hash is modified. The comparator is
not the byte-identical original and does not establish an original hash.
Only a current diagnostic JSON under .work/ap18-archive is written.
"""
import ast
from copy import deepcopy
from datetime import timedelta
import hashlib
import io
import json
from pathlib import Path
import re
import runpy
import sys
import xml.etree.ElementTree as ET
import zipfile
from openpyxl.formula.translate import Translator

ROOT = Path(__file__).resolve().parents[1]
ADAPTER_FILE = ROOT / 'scripts/finalize-ap18b-03.py'
AUDIT = runpy.run_path(str(ROOT / 'scripts/check-ap18b-03-workbook.py'))
Book, NS = AUDIT['Book'], AUDIT['NS']
V09_SHA = 'a245080126586f1104b459ff5e225808e2d78578d24a2ce619d7d31d2d27ed12'
V10_SHA = 'bba15a8f6e7b6956662d00bb1b7efb51903ad19aa2b2d1dc29c42a7fd9361e4e'
QA09, QA10 = ROOT / '.work/ap18b-03/batch', ROOT / '.work/ap18b-04/batch'
CHECKS09, CHECKS10 = [json.loads((path / 'checks.json').read_text()) for path in (QA09, QA10)]
TARGET = Path(CHECKS10['file'])
assert CHECKS10['sourceSha256'] == V09_SHA
assert Path(CHECKS09['file']) == Path(CHECKS10['source'])


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reconstruct(mode, predecessor=None):
    """Keep native adapter pre-write checks, with explicit comparator exception.

    For V0.10 only, the old file-hash assertion is removed because the
    historical V0.9 comparator is reconstructed, not recovered as bytes.
    Actual source integrity remains checked separately before/after this run.
    All workbook structural and permitted-edit checks remain in force.
    """
    env = runpy.run_path(str(ADAPTER_FILE))
    tree = ast.parse(ADAPTER_FILE.read_text())
    main = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'main')
    stop = next(i for i, n in enumerate(main.body)
                if isinstance(n, ast.Expr) and isinstance(n.value, ast.Call)
                and isinstance(n.value.func, ast.Attribute) and n.value.func.attr == 'mkdir')
    main.body = main.body[:stop]
    if predecessor is not None:
        matches = [n for n in main.body if isinstance(n, ast.Expr)
                   and isinstance(n.value, ast.Call)
                   and any(isinstance(a, ast.Constant) and a.value == 'source SHA mismatch'
                           for a in n.value.args)]
        assert len(matches) == 1
        main.body.remove(matches[0])
        original_read = env['read']

        def read(path):
            return deepcopy(predecessor) if Path(path) == Path(CHECKS10['source']) else original_read(path)

        env['read'] = read
    main.body.append(ast.Return(value=ast.Name(id='parts', ctx=ast.Load())))
    module = ast.fix_missing_locations(ast.Module(body=[main], type_ignores=[]))
    exec(compile(module, str(ADAPTER_FILE) + ':diagnostic-read-only-prefix', 'exec'), env)
    previous = sys.argv
    try:
        sys.argv = [str(ADAPTER_FILE), mode]
        return env['main']()
    finally:
        sys.argv = previous


def as_book(parts):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        for key, value in parts.items():
            archive.writestr(key, value)
    buffer.seek(0)
    return Book(buffer)


def expand_shared(book):
    count = 0
    for cells in book.cells.values():
        origins = {}
        for address, cell in cells.items():
            formula = cell.find('s:f', NS)
            if formula is not None and formula.get('t') == 'shared' and formula.text:
                origins[formula.get('si')] = (address, formula.text)
        for address, cell in cells.items():
            formula = cell.find('s:f', NS)
            if formula is not None and formula.get('t') == 'shared' and not formula.text:
                origin, text = origins[formula.get('si')]
                formula.text = Translator('=' + text, origin=origin).translate_formula(address)[1:]
                count += 1
    return count


def normal_formula(value, names):
    if value is not None:
        for name in names:
            assert not re.search(r"[ !'\[\]]", name)
            value = value.replace("'" + name + "'!", name + '!')
    return value


def normal_xml(element, names):
    if element is None:
        return None
    attrs = {k: v for k, v in element.attrib.items()
             if not k.startswith('{http://schemas.microsoft.com/office/spreadsheetml/')
             and not k.startswith('{http://schemas.openxmlformats.org/markup-compatibility/')}
    defaults = {
        'table': {'headerRowCount': '1', 'totalsRowCount': '0'},
        'sheetProtection': {'selectLockedCells': '0', 'selectUnlockedCells': '0'},
        'dataValidation': {'operator': 'between', 'errorStyle': 'stop'},
    }
    tag = element.tag.split('}')[-1]
    for k, v in defaults.get(tag, {}).items():
        attrs.setdefault(k, v)
    text = element.text or ''
    if tag in ('formula', 'formula1', 'formula2', 'f'):
        text = normal_formula(text, names)
    return [element.tag, sorted(attrs.items()), text, [normal_xml(c, names) for c in element]]


def normal_feature(book, element):
    element = deepcopy(element)
    styles = book.xml['xl/styles.xml'].find('s:dxfs', NS)
    for rule in element.findall('s:cfRule', NS):
        if 'dxfId' in rule.attrib:
            style = styles[int(rule.attrib.pop('dxfId'))]
            rule.set('diagnosticResolvedDxf', json.dumps(normal_xml(style, book.paths), ensure_ascii=False))
    return normal_xml(element, book.paths)


def hyperlinks(book, name):
    path = AUDIT['relationship_part'](book.paths[name])
    rels = {r.get('Id'): r for r in book.xml.get(path, [])}
    result = []
    for link in book.sheets[name].findall('s:hyperlinks/s:hyperlink', NS):
        rid = link.get('{' + AUDIT['R'] + '}id')
        result.append([link.get('ref'), link.get('location'),
                       rels[rid].get('Target') if rid else None,
                       rels[rid].get('TargetMode') if rid else None])
    return sorted(result, key=lambda row: row[0])


def main():
    sources = [Path(CHECKS09['source']), Path(CHECKS09['file']), TARGET,
               QA09 / 'checks.json', QA10 / 'checks.json',
               QA09 / 'artifact-export.xlsx', QA10 / 'artifact-export.xlsx']
    before = {path: digest(path) for path in sources}
    expected = as_book(reconstruct('--batch-04', reconstruct('--batch')))
    current = Book(TARGET)
    shared = expand_shared(current)
    assert set(expected.paths) == set(current.paths)
    values, formulas, locks, missing, extra, raw_formula_count = [], [], [], [], [], 0
    cells_count = formula_count = 0
    for name, cells in expected.cells.items():
        for address in cells:
            cells_count += 1
            a, b = expected.value(name, address), current.value(name, address)
            fa, fb = expected.formula(name, address), current.formula(name, address)
            if a != b and not (a in (None, '') and b in (None, '')):
                values.append([name, address, a, b])
            formula_count += fa is not None
            raw_formula_count += fa != fb
            if normal_formula(fa, expected.paths) != normal_formula(fb, expected.paths):
                formulas.append([name, address, fa, fb])
            if address not in current.cells[name] and (a not in (None, '') or fa is not None):
                missing.append([name, address])
            if address in current.cells[name] and expected.locked(name, address) != current.locked(name, address):
                locks.append([name, address])
        extra.extend([name, address] for address in current.cells[name].keys() - cells.keys()
                     if current.value(name, address) not in (None, '') or current.formula(name, address))
    tables = [name for name, (_, table) in expected.tables.items()
              if normal_xml(table, expected.paths) != normal_xml(current.tables[name][1], current.paths)]
    features = {}
    for tag in ('sheetProtection', 'mergeCells', 'autoFilter', 'dataValidations', 'conditionalFormatting'):
        features[tag] = [name for name in expected.paths
                         if sorted([normal_feature(expected, e) for e in expected.sheets[name].findall('s:' + tag, NS)], key=str)
                         != sorted([normal_feature(current, e) for e in current.sheets[name].findall('s:' + tag, NS)], key=str)]
    priorities, cf_other = [], []
    for name in features['conditionalFormatting']:
        old = {e.get('sqref'): e for e in expected.sheets[name].findall('s:conditionalFormatting', NS)}
        new = {e.get('sqref'): e for e in current.sheets[name].findall('s:conditionalFormatting', NS)}
        for scope in old.keys() | new.keys():
            if scope not in old or scope not in new:
                cf_other.append([name, scope, 'range added or removed'])
                continue
            a, b = deepcopy(old[scope]), deepcopy(new[scope])
            if normal_feature(expected, a) == normal_feature(current, b):
                continue
            pa = [r.attrib.pop('priority', None) for r in a.findall('s:cfRule', NS)]
            pb = [r.attrib.pop('priority', None) for r in b.findall('s:cfRule', NS)]
            disjoint = all(not (AUDIT['addresses'](scope) & AUDIT['addresses'](other)) for other in old if other != scope)
            if disjoint and normal_feature(expected, a) == normal_feature(current, b):
                priorities.append({'sheet': name, 'range': scope, 'before': pa, 'after': pb, 'otherRulesDoNotOverlap': True})
            else:
                cf_other.append([name, scope, 'not explained by disjoint priority renumbering'])
    links = [name for name in expected.paths if hyperlinks(expected, name) != hyperlinks(current, name)]
    date_checks, date_error, cache_checks = None, None, 0
    independent_inputs = [change for change in values
                          if expected.formula(change[0], change[1]) is None]
    unexplained_caches = []
    selected_year = current.value('Übersicht', 'B4')
    if not formulas:
        checked = deepcopy(current)
        for name, cells in expected.cells.items():
            for address in cells:
                formula = expected.formula(name, address)
                if formula is not None:
                    checked.cells[name][address].find('s:f', NS).text = formula
        rule_rows = {checked.value('Feiertagsregeln', f'A{r}'): r
                     for r in range(7, checked.end('Feiertagsregeln') + 1)}
        try:
            assert type(selected_year) is int and 2026 <= selected_year <= 2028
            date_checks = 0
            for year in (2026, 2027, 2028):
                evaluator = AUDIT['Evaluator'](checked, year)
                for key, row in rule_rows.items():
                    raw, valid = AUDIT['parameter_date'](checked, row, year)
                    for column, wanted in [('W', raw), ('X', valid)]:
                        assert evaluator.cell('Feiertagsregeln', f'{column}{row}') == wanted, (key, year, column)
                        date_checks += 1
            selected = AUDIT['Evaluator'](checked, selected_year)
            for name, address, _, value in values:
                if expected.formula(name, address) is None:
                    continue
                if name == 'Feiertagsregeln' and re.fullmatch(r'(W|X|AD)\d+', address):
                    wanted = selected.cell(name, address)
                elif name == 'Feiertagskalender' and re.fullmatch(r'[AB]\d+', address):
                    row = int(address[1:])
                    rule_row = rule_rows[checked.value(name, f'K{row}')]
                    serial = AUDIT['parameter_date'](checked, rule_row, selected_year)[1]
                    assert type(serial) is int
                    if address.startswith('A'):
                        wanted = serial
                    else:
                        day = AUDIT['EPOCH'] + timedelta(days=serial)
                        wanted = ['Montag', 'Dienstag', 'Mittwoch', 'Donnerstag', 'Freitag', 'Samstag', 'Sonntag'][day.weekday()]
                        if checked.value('Feiertagsregeln', f'AA{rule_row}') == 'Ab 12.00 Uhr':
                            wanted += ', ab 12.00 Uhr'
                else:
                    unexplained_caches.append([name, address, 'not covered by declared year-dependent cache checks'])
                    continue
                if value != wanted:
                    unexplained_caches.append([name, address, value, wanted])
                cache_checks += 1
        except Exception as error:
            date_error = str(error)
    assert all(digest(path) == value for path, value in before.items())
    report = {
        'readOnly': True, 'workbooksMutated': False, 'trustedHashesChanged': False,
        'historicalReportsOverwritten': False,
        'target': str(TARGET.relative_to(ROOT)), 'targetSha256': before[TARGET],
        'expectedHistoricalSha256': V10_SHA,
        'comparisonBasis': 'V0.8 reference, retained AP18B-03 and AP18B-04 author exports, declared patches, native adapter pre-write checks, replayed entirely in memory.',
        'byteIdenticalOriginalLocated': False,
        'predecessorHashGateNotAssertedForDiagnosticComparator': True,
        'inputDigests': {str(path.relative_to(ROOT)): value for path, value in before.items()},
        'comparedCells': cells_count, 'formulaCells': formula_count,
        'sharedFormulaFollowersExpandedReadOnly': shared,
        'valueChanges': values, 'rawFormulaChanges': raw_formula_count,
        'independentInputChanges': independent_inputs,
        'selectedYearBefore': expected.value('Übersicht', 'B4'),
        'selectedYearAfter': selected_year,
        'changedFormulaCachesIndependentlyChecked': cache_checks,
        'unexplainedChangedFormulaCaches': unexplained_caches,
        'semanticFormulaChangesAfterOptionalSheetQuoteNormalisation': formulas,
        'missingNonEmptyCells': missing, 'extraNonEmptyCells': extra,
        'cellProtectionChanges': locks,
        'nativeTableDifferencesIgnoringExcelRevisionIds': tables,
        'nativeFeatureDifferencesIgnoringExcelRevisionIds': features,
        'cfPriorityOnlyChangesOnDisjointRanges': priorities,
        'otherConditionalFormattingChanges': cf_other,
        'resolvedHyperlinkChanges': links,
        'dateFormulaComparisons2026To2028': date_checks, 'dateCheckError': date_error,
        'addedPackageParts': sorted(current.raw.keys() - expected.raw.keys()),
        'removedPackageParts': sorted(expected.raw.keys() - current.raw.keys()),
        'rawChangedPackageParts': sorted(k for k in current.raw.keys() & expected.raw.keys() if current.raw[k] != expected.raw[k]),
        'limitations': [
            'No byte-identical original located. A reconstructed comparator cannot prove the original XLSX hash.',
            'The V0.10 predecessor-hash assertion is deliberately not used for the reconstructed in-memory V0.9 comparator. No production or historical gate is changed.',
            'OOXML defaults, optional quotes around verified sheet names, Excel revision UIDs, and reordered differential-style pools are normalised for feature comparison.',
            'Excel view state, application metadata, and raw style/XML serialisation are not claimed byte-identical.',
            'No legal approval, archival acceptance, or product release is established by this diagnostic.',
        ],
    }
    output = ROOT / '.work/ap18-archive/v10-drift-diagnostic.json'
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: v for k, v in report.items()
                      if k not in ('inputDigests', 'rawChangedPackageParts', 'valueChanges')}
                     | {'valueChanges': len(values)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
